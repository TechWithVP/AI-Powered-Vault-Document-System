"""DOCX text extraction."""

from __future__ import annotations

import io
from zipfile import BadZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.core.constants import FileType
from app.core.exceptions import TextExtractionError
from app.processing.extractors.base import (
    ExtractedDocument,
    ExtractedPage,
    TextExtractor,
)

#: Word has no fixed pagination, so blocks are grouped into synthetic pages of
#: roughly this many characters to keep citations usefully specific.
SYNTHETIC_PAGE_CHARS = 3000


class DocxExtractor(TextExtractor):
    """Extracts text from Word documents, including table contents."""

    supported_types = frozenset({FileType.DOCX})

    def extract(self, data: bytes, *, filename: str) -> ExtractedDocument:
        try:
            document = DocxDocument(io.BytesIO(data))
        except (PackageNotFoundError, BadZipFile) as exc:
            # A .docx is a zip archive; non-zip bytes mean the file is either a
            # legacy .doc or simply mislabelled.
            raise TextExtractionError(
                "The file is not a valid .docx document. Legacy .doc files are "
                "not supported - please save it as .docx and retry."
            ) from exc
        except Exception as exc:
            raise TextExtractionError(f"Could not read the document: {exc}") from exc

        blocks = self._collect_blocks(document)
        pages = self._paginate(blocks)
        metadata = self._read_metadata(document)
        return ExtractedDocument(pages=pages, metadata=metadata).require_text()

    @staticmethod
    def _collect_blocks(document: object) -> list[str]:
        blocks = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return blocks

    @staticmethod
    def _paginate(blocks: list[str]) -> list[ExtractedPage]:
        pages: list[ExtractedPage] = []
        buffer: list[str] = []
        length = 0

        for block in blocks:
            buffer.append(block)
            length += len(block)
            if length >= SYNTHETIC_PAGE_CHARS:
                pages.append(
                    ExtractedPage(number=len(pages) + 1, text="\n".join(buffer))
                )
                buffer, length = [], 0

        if buffer:
            pages.append(ExtractedPage(number=len(pages) + 1, text="\n".join(buffer)))
        return pages

    @staticmethod
    def _read_metadata(document: object) -> dict[str, str]:
        properties = document.core_properties
        candidates = {
            "title": properties.title,
            "author": properties.author,
            "subject": properties.subject,
        }
        return {key: str(value)[:500] for key, value in candidates.items() if value}
